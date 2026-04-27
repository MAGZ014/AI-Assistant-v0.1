import os
import shutil
from datetime import datetime
from config import DOCS_FOLDER

os.makedirs(DOCS_FOLDER, exist_ok=True)

# ─────────────────────────────────────────
#  LEER ARCHIVOS (txt, md, docx)
# ─────────────────────────────────────────

def read_file(filename: str, max_chars: int = 3000) -> str:
    """
    Lee cualquier archivo de texto.
    max_chars evita saturar el contexto del LLM con archivos muy largos.
    """
    path = os.path.join(DOCS_FOLDER, filename)

    if not os.path.exists(path):
        return f"No encontré '{filename}' en {DOCS_FOLDER}."

    try:
        ext = filename.lower().split(".")[-1]

        if ext == "docx":
            from docx import Document
            doc = Document(path)
            text = "\n".join([p.text for p in doc.paragraphs if p.text.strip()])

        elif ext == "xlsx":
            import openpyxl
            wb = openpyxl.load_workbook(path)
            ws = wb.active
            rows = []
            for row in ws.iter_rows(values_only=True):
                rows.append(" | ".join([str(c) if c is not None else "" for c in row]))
            text = "\n".join(rows)

        else:  # txt, md y cualquier texto plano
            with open(path, "r", encoding="utf-8") as f:
                text = f.read()

        # Truncar si es muy largo
        if len(text) > max_chars:
            text = text[:max_chars] + f"\n\n... [truncado, {len(text) - max_chars} caracteres restantes]"

        return f"Contenido de '{filename}':\n{text}"

    except Exception as e:
        return f"Error leyendo '{filename}': {str(e)}"


# ─────────────────────────────────────────
#  LISTAR ARCHIVOS
# ─────────────────────────────────────────

def list_files(folder: str = None) -> str:
    """
    Lista archivos en DOCS_FOLDER o en una subcarpeta específica.
    Muestra nombre, tamaño y fecha de modificación.
    """
    target = os.path.join(DOCS_FOLDER, folder) if folder else DOCS_FOLDER

    if not os.path.exists(target):
        return f"La carpeta '{target}' no existe."

    try:
        entries = os.listdir(target)
        if not entries:
            return f"La carpeta está vacía."

        lines = []
        for entry in sorted(entries):
            full_path = os.path.join(target, entry)
            is_dir = os.path.isdir(full_path)

            if is_dir:
                lines.append(f"📁 {entry}/")
            else:
                size  = os.path.getsize(full_path)
                mtime = os.path.getmtime(full_path)
                date  = datetime.fromtimestamp(mtime).strftime("%d/%m/%Y %H:%M")
                size_str = f"{size}B" if size < 1024 else f"{size // 1024}KB"
                lines.append(f"📄 {entry}  [{size_str}  {date}]")

        return f"Contenido de '{target}':\n" + "\n".join(lines)

    except Exception as e:
        return f"Error listando archivos: {str(e)}"


# ─────────────────────────────────────────
#  MOVER Y RENOMBRAR
# ─────────────────────────────────────────

def move_file(filename: str, destination: str) -> str:
    """Mueve un archivo a otra carpeta."""
    source = os.path.join(DOCS_FOLDER, filename)

    if not os.path.exists(source):
        return f"No encontré '{filename}'."

    if not os.path.isabs(destination):
        destination = os.path.join(DOCS_FOLDER, destination)

    try:
        os.makedirs(destination, exist_ok=True)
        shutil.move(source, os.path.join(destination, filename))
        return f"'{filename}' movido a '{destination}'."
    except Exception as e:
        return f"Error moviendo '{filename}': {str(e)}"


def rename_file(filename: str, new_name: str) -> str:
    """Renombra un archivo."""
    old_path = os.path.join(DOCS_FOLDER, filename)
    new_path = os.path.join(DOCS_FOLDER, new_name)

    if not os.path.exists(old_path):
        return f"No encontré '{filename}'."

    try:
        os.rename(old_path, new_path)
        return f"'{filename}' renombrado a '{new_name}'."
    except Exception as e:
        return f"Error renombrando: {str(e)}"


# ─────────────────────────────────────────
#  COPIAR
# ─────────────────────────────────────────

def copy_file(filename: str, new_name: str) -> str:
    """Copia un archivo con un nuevo nombre dentro de DOCS_FOLDER."""
    source = os.path.join(DOCS_FOLDER, filename)
    dest   = os.path.join(DOCS_FOLDER, new_name)

    if not os.path.exists(source):
        return f"No encontré '{filename}'."

    try:
        shutil.copy2(source, dest)
        return f"'{filename}' copiado como '{new_name}'."
    except Exception as e:
        return f"Error copiando: {str(e)}"


# ─────────────────────────────────────────
#  ELIMINAR
# ─────────────────────────────────────────

def delete_file(filename: str) -> str:
    """Elimina un archivo."""
    path = os.path.join(DOCS_FOLDER, filename)

    if not os.path.exists(path):
        return f"No encontré '{filename}'."

    try:
        os.remove(path)
        return f"'{filename}' eliminado."
    except Exception as e:
        return f"Error eliminando: {str(e)}"


# ─────────────────────────────────────────
#  BUSCAR
# ─────────────────────────────────────────

def search_files(query: str) -> str:
    """
    Busca archivos por nombre o extensión en DOCS_FOLDER y subcarpetas.
    Ejemplo: search_files("receta") o search_files(".docx")
    """
    try:
        matches = []
        for root, dirs, files in os.walk(DOCS_FOLDER):
            for f in files:
                if query.lower() in f.lower():
                    rel_path = os.path.relpath(os.path.join(root, f), DOCS_FOLDER)
                    matches.append(rel_path)

        if not matches:
            return f"No encontré archivos que coincidan con '{query}'."

        return f"Archivos encontrados para '{query}':\n" + "\n".join([f"• {m}" for m in matches])

    except Exception as e:
        return f"Error buscando: {str(e)}"


# ─────────────────────────────────────────
#  INFO DE ARCHIVO
# ─────────────────────────────────────────

def file_info(filename: str) -> str:
    """Muestra información detallada de un archivo."""
    path = os.path.join(DOCS_FOLDER, filename)

    if not os.path.exists(path):
        return f"No encontré '{filename}'."

    try:
        size  = os.path.getsize(path)
        mtime = datetime.fromtimestamp(os.path.getmtime(path)).strftime("%d/%m/%Y %H:%M")
        ctime = datetime.fromtimestamp(os.path.getctime(path)).strftime("%d/%m/%Y %H:%M")
        ext   = filename.split(".")[-1].upper() if "." in filename else "Sin extensión"

        size_str = f"{size} bytes" if size < 1024 else f"{size // 1024} KB"

        return (
            f"Archivo: {filename}\n"
            f"Tipo: {ext}\n"
            f"Tamaño: {size_str}\n"
            f"Creado: {ctime}\n"
            f"Modificado: {mtime}"
        )
    except Exception as e:
        return f"Error obteniendo info: {str(e)}"


# ─────────────────────────────────────────
#  APPEND
# ─────────────────────────────────────────

def append_to_file(filename: str, content: str) -> str:
    """Agrega contenido al final de un archivo txt o md existente."""
    path = os.path.join(DOCS_FOLDER, filename)

    if not os.path.exists(path):
        return f"No encontré '{filename}'. ¿Quieres crearlo primero?"

    try:
        with open(path, "a", encoding="utf-8") as f:
            f.write(f"\n{content}")
        return f"Contenido agregado a '{filename}'."
    except Exception as e:
        return f"Error agregando contenido: {str(e)}"